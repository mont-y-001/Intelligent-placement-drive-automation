from pathlib import Path
from collections import Counter

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).parent
DATA = ROOT / "Data"
OUT = ROOT / "Intelligent_Placement_System_Evaluation.pptx"

NAVY = RGBColor(12, 30, 50)
INK = RGBColor(24, 35, 48)
TEAL = RGBColor(0, 150, 148)
ORANGE = RGBColor(245, 143, 56)
PALE = RGBColor(241, 247, 246)
WHITE = RGBColor(255, 255, 255)
MUTED = RGBColor(91, 108, 122)
RED = RGBColor(185, 61, 61)


def add_text(slide, text, x, y, w, h, size=20, color=INK, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_bullets(slide, items, x, y, w, h, size=19, color=INK):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.04)
    for index, item in enumerate(items):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.space_after = Pt(10)
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.font.bold = False
    return box


def add_header(slide, title, section):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.22))
    band.fill.solid()
    band.fill.fore_color.rgb = TEAL
    band.line.fill.background()
    add_text(slide, section.upper(), 0.65, 0.52, 2.4, 0.25, 10, TEAL, True)
    add_text(slide, title, 0.65, 0.82, 12.0, 0.6, 28, NAVY, True)
    add_text(slide, "Intelligent Placement System | Evaluation Presentation", 0.65, 7.12, 7.0, 0.2, 9, MUTED)


def add_card(slide, title, body, x, y, w, h, accent=TEAL):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = PALE
    shape.line.color.rgb = RGBColor(218, 232, 231)
    strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h))
    strip.fill.solid()
    strip.fill.fore_color.rgb = accent
    strip.line.fill.background()
    add_text(slide, title, x + 0.25, y + 0.18, w - 0.45, 0.35, 16, NAVY, True)
    add_text(slide, body, x + 0.25, y + 0.62, w - 0.45, h - 0.78, 13, INK)


def add_footer_number(slide, number):
    add_text(slide, str(number), 12.35, 7.08, 0.35, 0.25, 10, TEAL, True, PP_ALIGN.RIGHT)


def add_flow_step(slide, number, title, body, x, y, w=2.15):
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(0.5), Inches(0.5))
    circle.fill.solid()
    circle.fill.fore_color.rgb = TEAL
    circle.line.fill.background()
    add_text(slide, str(number), x, y + 0.02, 0.5, 0.42, 15, WHITE, True, PP_ALIGN.CENTER)
    add_text(slide, title, x, y + 0.7, w, 0.35, 14, NAVY, True)
    add_text(slide, body, x, y + 1.12, w, 1.1, 12, INK)


def build_deck():
    students = pd.read_excel(DATA / "student_data_100.xlsx")
    skill_counts = Counter()
    for value in students["Skills"].fillna(""):
        skill_counts.update(skill.strip() for skill in str(value).split(",") if skill.strip())
    top_skills = skill_counts.most_common(5)
    current_required = ["Python", "HTML", "CSS", "JavaScript", "Flask/Django", "React/Angular/Vue", "SQL", "Git", "Problem-solving"]
    matched_counts = {skill: int(skill_counts.get(skill, 0)) for skill in current_required}
    skill_only_eligible = students["Skills"].fillna("").apply(
        lambda value: all(skill.lower() in [s.strip().lower() for s in str(value).split(",")] for skill in ["Python", "SQL"])
    ).sum()

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.28), Inches(7.5))
    accent.fill.solid(); accent.fill.fore_color.rgb = ORANGE; accent.line.fill.background()
    add_text(slide, "INTELLIGENT PLACEMENT SYSTEM", 0.85, 0.9, 7.8, 0.35, 15, RGBColor(103, 220, 210), True)
    add_text(slide, "From Job Description\nto Eligible Candidates", 0.82, 1.55, 8.3, 1.5, 36, WHITE, True)
    add_text(slide, "AI-assisted placement screening using PDF retrieval, structured JD analysis, and student skill matching", 0.88, 3.45, 7.4, 0.75, 18, RGBColor(224, 237, 239))
    add_text(slide, "MSE-1 Robotic Agentic Automation (RAA) | CA212E", 0.88, 6.55, 7.5, 0.3, 13, WHITE, True)
    add_text(slide, "Evaluation presentation", 0.88, 6.9, 5.5, 0.25, 11, RGBColor(170, 191, 199))
    for index, label in enumerate(["INGEST", "ANALYZE", "MATCH"]):
        add_card(slide, label, ["PDF + Excel", "Gemini + Pydantic", "Eligibility shortlist"][index], 9.2, 1.45 + index * 1.35, 2.65, 0.95, [TEAL, ORANGE, RGBColor(80, 165, 218)][index])

    slide = prs.slides.add_slide(blank); add_header(slide, "Evaluation alignment at a glance", "01 | Evaluation map"); add_footer_number(slide, 2)
    add_card(slide, "Problem identification & requirement analysis", "Manual screening is slow and inconsistent when a JD and a large student sheet must be compared.", 0.8, 1.8, 5.7, 1.35, TEAL)
    add_card(slide, "Process analysis & RPA solution design", "A repeatable pipeline turns unstructured JD text into structured criteria, then screens student records.", 6.85, 1.8, 5.7, 1.35, ORANGE)
    add_card(slide, "Proposed solution, novelty & feasibility", "Agent-based analysis, vector retrieval, and a transparent matching rule create an explainable prototype.", 0.8, 3.55, 5.7, 1.35, RGBColor(80, 165, 218))
    add_card(slide, "Evidence used in this deck", "Local JD PDF, 100-row Excel dataset, source implementation, and the current prototype behavior.", 6.85, 3.55, 5.7, 1.35, RGBColor(120, 157, 91))
    add_text(slide, "The deck separates what is implemented today from what is recommended for the next iteration.", 0.85, 5.65, 11.8, 0.45, 18, NAVY, True, PP_ALIGN.CENTER)

    slide = prs.slides.add_slide(blank); add_header(slide, "The placement-screening problem", "02 | Problem"); add_footer_number(slide, 3)
    add_bullets(slide, [
        "Job descriptions are semi-structured documents with requirements mixed with responsibilities and optional skills.",
        "Student information is stored separately in tabular form, creating a comparison and normalization problem.",
        "Manual review does not scale to 100+ candidates and makes decisions harder to audit.",
        "The target outcome is a shortlist with a clear reason for inclusion or exclusion."
    ], 0.9, 1.65, 7.25, 3.7, 19)
    add_card(slide, "Input reality", "JD: PDF\nCandidates: Excel\nDecision: eligibility shortlist", 8.65, 1.85, 3.75, 2.2, ORANGE)
    add_text(slide, "Success measure", 8.75, 4.55, 3.4, 0.35, 15, TEAL, True)
    add_text(slide, "Reduce repetitive screening while preserving traceability from JD evidence to candidate decision.", 8.75, 4.95, 3.35, 1.15, 18, NAVY, True)

    slide = prs.slides.add_slide(blank); add_header(slide, "System architecture", "03 | Process analysis"); add_footer_number(slide, 4)
    add_flow_step(slide, 1, "PDF Loader", "PyPDFLoader reads the job description and returns document pages.", 0.9, 1.75)
    add_flow_step(slide, 2, "Text Splitter", "The JD is split into chunks for focused retrieval.", 3.25, 1.75)
    add_flow_step(slide, 3, "Vector Store", "Hugging Face embeddings and FAISS index the JD chunks.", 5.6, 1.75)
    add_flow_step(slide, 4, "JD Agent", "Gemini returns a typed JDRequirements object.", 7.95, 1.75)
    add_flow_step(slide, 5, "Eligibility Agent", "Student skill lists are compared with required skills.", 10.3, 1.75)
    for x in [2.95, 5.3, 7.65, 10.0]:
        line = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(2.0), Inches(0.32), Inches(0.28))
        line.fill.solid(); line.fill.fore_color.rgb = ORANGE; line.line.fill.background()
    add_text(slide, "LangChain orchestration", 1.0, 5.05, 3.2, 0.35, 16, TEAL, True)
    add_text(slide, "PDFLoader -> TextSplitter -> Embedding -> Retriever -> JDAgent -> EligibilityAgent", 1.0, 5.48, 11.2, 0.48, 19, NAVY, True, PP_ALIGN.CENTER)

    slide = prs.slides.add_slide(blank); add_header(slide, "Workflow: from source documents to shortlist", "04 | Solution design"); add_footer_number(slide, 5)
    steps = [
        ("1", "Load", "Read the JD PDF and student workbook."),
        ("2", "Retrieve", "Fetch the most relevant JD context."),
        ("3", "Structure", "Extract role, education, skills, and experience."),
        ("4", "Match", "Compare mandatory skills with each student."),
        ("5", "Report", "Return candidate details for review."),
    ]
    for i, (num, title, body) in enumerate(steps):
        x = 0.85 + i * 2.45
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(1.85), Inches(0.72), Inches(0.72))
        circle.fill.solid(); circle.fill.fore_color.rgb = [TEAL, ORANGE, RGBColor(80,165,218), RGBColor(120,157,91), NAVY][i]; circle.line.fill.background()
        add_text(slide, num, x, 1.94, 0.72, 0.4, 20, WHITE, True, PP_ALIGN.CENTER)
        add_text(slide, title, x - 0.18, 2.82, 1.1, 0.35, 16, NAVY, True, PP_ALIGN.CENTER)
        add_text(slide, body, x - 0.35, 3.3, 1.55, 1.25, 13, INK, False, PP_ALIGN.CENTER)
        if i < 4:
            arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + 0.9), Inches(2.06), Inches(1.15), Inches(0.3))
            arrow.fill.solid(); arrow.fill.fore_color.rgb = RGBColor(210, 223, 224); arrow.line.fill.background()
    add_card(slide, "Design principle", "Keep extraction and matching separate: the JD agent interprets the document; the eligibility agent applies a visible rule to student rows.", 1.55, 5.35, 10.2, 0.95, ORANGE)

    slide = prs.slides.add_slide(blank); add_header(slide, "Job-description analysis is structured, not free-form", "05 | Technical novelty"); add_footer_number(slide, 6)
    add_text(slide, "Source JD: Python Full Stack Developer Fresher", 0.9, 1.55, 6.3, 0.4, 20, NAVY, True)
    add_bullets(slide, [
        "Role: Python Full Stack Developer Fresher",
        "Core work: backend systems, frontend applications, APIs, databases",
        "Required skills: Python, HTML, CSS, JavaScript, backend framework, frontend framework, SQL, Git",
        "Good to have: internship/project experience, cloud, Docker or CI/CD basics"
    ], 1.0, 2.1, 6.3, 3.3, 16)
    add_card(slide, "Typed output", "JDRequirements\n\njob_role\nrequired_qualification\nminimum_10th_percentage\nminimum_12th_percentage\nminimum_graduation_percentage\nrequired_technical_skills\nrequired_experience", 8.05, 1.55, 3.95, 3.85, TEAL)
    add_text(slide, "Pydantic validation makes downstream matching predictable and easier to test.", 8.1, 5.8, 3.8, 0.75, 15, NAVY, True, PP_ALIGN.CENTER)

    slide = prs.slides.add_slide(blank); add_header(slide, "Candidate data and current matching result", "06 | Evidence"); add_footer_number(slide, 7)
    add_text(slide, f"Dataset size: {len(students)} student records", 0.9, 1.45, 4.8, 0.35, 18, NAVY, True)
    add_text(slide, f"Current skill-only example (Python + SQL): {skill_only_eligible} candidates", 0.9, 1.9, 6.8, 0.35, 16, TEAL, True)
    chart_x, chart_y, chart_w, chart_h = 0.95, 2.75, 7.0, 2.85
    max_count = max(count for _, count in top_skills) or 1
    for i, (skill, count) in enumerate(top_skills):
        y = chart_y + i * 0.52
        add_text(slide, skill, chart_x, y, 1.45, 0.3, 13, INK, True)
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(chart_x + 1.55), Inches(y + 0.05), Inches(4.65 * count / max_count), Inches(0.24))
        bar.fill.solid(); bar.fill.fore_color.rgb = TEAL; bar.line.fill.background()
        add_text(slide, str(count), chart_x + 6.35, y - 0.01, 0.45, 0.3, 13, NAVY, True, PP_ALIGN.RIGHT)
    add_card(slide, "Sample record", "Aarav Sharma\nMCA2026001\n84.82% | 60.88% | 66.08%\nSkills: Python, ML", 8.5, 2.0, 3.2, 2.35, ORANGE)
    add_text(slide, "The workbook is ready for batch screening; the current agent returns matching rows as candidate dictionaries.", 8.55, 4.75, 3.15, 1.1, 15, NAVY, True, PP_ALIGN.CENTER)

    slide = prs.slides.add_slide(blank); add_header(slide, "Evaluation scorecard", "07 | Assessment"); add_footer_number(slide, 8)
    rows = [
        ("Problem Identification & Requirement Analysis", "Clear manual-screening problem; source JD and Excel evidence", "Strong"),
        ("Process Analysis & RPA Solution Design", "Five-stage pipeline with separate analysis and matching agents", "Strong"),
        ("Proposed Solution", "PDF retrieval + structured extraction + deterministic skill matching", "Strong"),
        ("Novelty", "Agent separation and typed JD output improve explainability", "Good"),
        ("Feasibility", "Runs with existing Python/LangChain stack and local data", "Good"),
    ]
    y = 1.55
    for title, evidence, rating in rows:
        add_text(slide, title, 0.85, y, 4.3, 0.42, 14, NAVY, True)
        add_text(slide, evidence, 5.05, y, 5.75, 0.42, 13, INK)
        color = TEAL if rating == "Strong" else ORANGE
        tag = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(11.15), Inches(y - 0.02), Inches(1.15), Inches(0.42))
        tag.fill.solid(); tag.fill.fore_color.rgb = color; tag.line.fill.background()
        add_text(slide, rating, 11.15, y + 0.02, 1.15, 0.25, 11, WHITE, True, PP_ALIGN.CENTER)
        y += 0.86
    add_text(slide, "Rubric fit is strongest when the demo shows the complete path from JD evidence to candidate output.", 1.1, 6.1, 11.1, 0.4, 18, NAVY, True, PP_ALIGN.CENTER)

    slide = prs.slides.add_slide(blank); add_header(slide, "Current limitations and next improvements", "08 | Rigor"); add_footer_number(slide, 9)
    add_card(slide, "Implemented now", "PDF retrieval\nTyped JD extraction\nSkill-based eligibility\nCandidate details printed for review", 0.9, 1.55, 3.6, 2.6, TEAL)
    add_card(slide, "Important limitation", "The JD agent extracts academic thresholds, but EligibilityAgent currently checks only required technical skills. Percentages are not yet used in the decision.", 4.85, 1.55, 3.6, 2.6, RED)
    add_card(slide, "Recommended next step", "Apply education thresholds, normalize synonyms such as Flask/Django, add match reasons, and export a ranked shortlist.", 8.8, 1.55, 3.6, 2.6, ORANGE)
    add_text(slide, "This distinction protects evaluation credibility: the presentation demonstrates the prototype without overstating its current behavior.", 1.2, 5.25, 10.95, 0.75, 22, NAVY, True, PP_ALIGN.CENTER)

    slide = prs.slides.add_slide(blank); add_header(slide, "Demonstration plan and conclusion", "09 | Close"); add_footer_number(slide, 10)
    add_bullets(slide, [
        "Open with the source JD and show the extracted requirements object.",
        "Show the 100-row student workbook and run the eligibility agent.",
        "Display two or three candidate records with their matched skills.",
        "Explain the current limitation and the next release improvement.",
        "Conclude: the system converts placement screening into a repeatable, explainable workflow."
    ], 1.0, 1.6, 7.2, 3.9, 18)
    add_card(slide, "Key takeaway", "A practical agentic pipeline for faster placement screening, grounded in real institutional data and designed for incremental improvement.", 8.75, 2.0, 3.4, 2.35, TEAL)
    add_text(slide, "Thank you", 0.9, 6.35, 11.5, 0.45, 25, NAVY, True, PP_ALIGN.CENTER)

    prs.save(OUT)
    print(f"Created {OUT}")


if __name__ == "__main__":
    build_deck()
